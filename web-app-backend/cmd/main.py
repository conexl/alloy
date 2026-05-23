from fastapi import FastAPI, Request, HTTPException, Depends, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
import asyncpg
import redis.asyncio as redis
import os
import asyncio
import json
from typing import List, Optional, AsyncGenerator
from pydantic import BaseModel
from datetime import datetime
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="Tender Web App API", version="1.1.0")

# CORS для VK Mini App
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Database and Redis
db_pool: asyncpg.Pool = None
redis_client: redis.Redis = None

# SSE клиенты: {vk_user_id: [queue1, queue2, ...]}
sse_clients: dict[int, list[asyncio.Queue]] = {}

@app.on_event("startup")
async def startup():
    global db_pool, redis_client
    dsn = os.getenv("DB_DSN", "postgres://postgres:postgres@postgres:5432/tender?sslmode=disable")
    db_pool = await asyncpg.create_pool(dsn)
    
    redis_url = os.getenv("REDIS_URL", "redis://redis:6379")
    redis_client = redis.from_url(redis_url, decode_responses=True)
    
    # Запускаем слушатель Redis Stream
    asyncio.create_task(redis_stream_listener())
    logger.info("Web app backend started")

@app.on_event("shutdown")
async def shutdown():
    await db_pool.close()
    await redis_client.close()

async def get_current_user(request: Request) -> int:
    vk_user_id = request.headers.get("X-VK-User-ID")
    if not vk_user_id:
        raise HTTPException(status_code=401, detail="VK User ID required")
    return int(vk_user_id)

# ========== SSE (Real-time updates) ==========

async def redis_stream_listener():
    """Слушает Redis Stream quote_received и рассылает SSE клиентам"""
    while True:
        try:
            # Читаем из Redis Stream
            messages = await redis_client.xread({"quote_received": "$"}, block=5000, count=10)
            
            for stream_name, stream_messages in messages:
                for msg_id, msg_data in stream_messages:
                    try:
                        data = json.loads(msg_data.get("data", "{}"))
                        item_id = data.get("item_id")
                        rfq_id = data.get("rfq_id")
                        
                        if item_id and rfq_id:
                            # Отправляем всем клиентам
                            await broadcast_sse({
                                "type": "quote_received",
                                "item_id": item_id,
                                "rfq_id": rfq_id,
                                "supplier_id": data.get("supplier_id"),
                                "price": data.get("price"),
                                "timestamp": datetime.now().isoformat(),
                            })
                    except Exception as e:
                        logger.error(f"Error processing Redis message: {e}")
                        
        except Exception as e:
            logger.error(f"Redis stream listener error: {e}")
            await asyncio.sleep(5)

async def broadcast_sse(data: dict):
    """Отправить событие всем подключённым клиентам"""
    event_json = json.dumps(data)
    for user_id, queues in list(sse_clients.items()):
        for queue in queues:
            try:
                await queue.put(event_json)
            except Exception:
                pass

@app.get("/api/sse/subscribe")
async def sse_subscribe(request: Request):
    """SSE endpoint для real-time обновлений"""
    # Получаем vk_user_id из query params (EventSource не поддерживает заголовки)
    vk_user_id_str = request.query_params.get("vk_user_id") or request.headers.get("X-VK-User-ID")
    if not vk_user_id_str:
        raise HTTPException(status_code=401, detail="VK User ID required")
    
    try:
        vk_user_id = int(vk_user_id_str)
    except ValueError:
        raise HTTPException(status_code=401, detail="Invalid VK User ID")
    
    queue = asyncio.Queue()
    
    if vk_user_id not in sse_clients:
        sse_clients[vk_user_id] = []
    sse_clients[vk_user_id].append(queue)
    
    async def event_generator() -> AsyncGenerator[str, None]:
        try:
            # Отправляем начальное сообщение
            yield f"data: {json.dumps({'type': 'connected'})}\n\n"
            
            while True:
                try:
                    event = await asyncio.wait_for(queue.get(), timeout=30)
                    yield f"data: {event}\n\n"
                except asyncio.TimeoutError:
                    # Keep-alive
                    yield f"data: {json.dumps({'type': 'ping'})}\n\n"
        except asyncio.CancelledError:
            pass
        finally:
            # Удаляем клиента
            if vk_user_id in sse_clients and queue in sse_clients[vk_user_id]:
                sse_clients[vk_user_id].remove(queue)
                if not sse_clients[vk_user_id]:
                    del sse_clients[vk_user_id]
    
    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        }
    )

# ========== Excel Parser ==========

def parse_excel_simple(content: bytes) -> tuple[list[dict], Optional[str], Optional[float]]:
    """Простой парсер Excel (без pandas, используем openpyxl)"""
    try:
        import openpyxl
        from io import BytesIO
        
        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb.active
        
        # Определяем заголовки
        headers = []
        header_row = None
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if any(cell for cell in row if cell):
                headers = [str(cell).strip().lower() if cell else "" for cell in row]
                header_row = row_idx
                break
        
        if not headers:
            raise ValueError("Не найдены заголовки")
        
        # Маппинг колонок
        col_map = {}
        for idx, h in enumerate(headers):
            if any(alias in h for alias in ["наименование", "name", "название", "товар"]):
                col_map["name"] = idx
            elif any(alias in h for alias in ["количество", "qty", "кол-во", "кол"]):
                col_map["qty"] = idx
            elif any(alias in h for alias in ["единица", "unit", "ед.изм"]):
                col_map["unit"] = idx
            elif any(alias in h for alias in ["срок", "deadline", "дата"]):
                col_map["deadline"] = idx
            elif any(alias in h for alias in ["город", "city", "место"]):
                col_map["location"] = idx
            elif any(alias in h for alias in ["бюджет", "budget", "сумма"]):
                col_map["budget"] = idx
            elif any(alias in h for alias in ["тех", "tech", "код"]):
                col_map["tech_code"] = idx
            elif any(alias in h for alias in ["гост", "gost", "ту", "стандарт"]):
                col_map["gost"] = idx
            elif any(alias in h for alias in ["материал", "material"]):
                col_map["material"] = idx
        
        if "name" not in col_map:
            raise ValueError("Колонка 'Наименование' не найдена")
        if "qty" not in col_map:
            raise ValueError("Колонка 'Количество' не найдена")
        
        items = []
        city = None
        budget = None
        
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row or not row[col_map["name"]]:
                continue
            
            name = str(row[col_map["name"]]).strip()
            if not name:
                continue
            
            # Количество
            qty_raw = row[col_map["qty"]]
            try:
                qty = float(qty_raw) if qty_raw else 1
            except (ValueError, TypeError):
                qty = 1
            
            # Единица
            unit = str(row[col_map.get("unit", -1)]).strip() if col_map.get("unit") and row[col_map["unit"]] else "шт"
            
            # Срок
            deadline = None
            if col_map.get("deadline") and row[col_map["deadline"]]:
                try:
                    if isinstance(row[col_map["deadline"]], datetime):
                        deadline = row[col_map["deadline"]]
                    else:
                        from dateutil import parser
                        deadline = parser.parse(str(row[col_map["deadline"]]))
                except:
                    pass
            
            # Город
            location = None
            if col_map.get("location") and row[col_map["location"]]:
                location = str(row[col_map["location"]]).strip()
                if location and not city:
                    city = location
            
            # Бюджет
            if col_map.get("budget") and row[col_map["budget"]]:
                try:
                    budget = float(row[col_map["budget"]])
                except:
                    pass
            
            item = {
                "name": name,
                "qty": qty,
                "unit": unit,
                "deadline": deadline,
                "location": location,
                "tech_code": str(row[col_map["tech_code"]]).strip() if col_map.get("tech_code") and row[col_map["tech_code"]] else None,
                "gost": str(row[col_map["gost"]]).strip() if col_map.get("gost") and row[col_map["gost"]] else None,
                "material": str(row[col_map["material"]]).strip() if col_map.get("material") and row[col_map["material"]] else None,
            }
            items.append(item)
        
        return items, city, budget
        
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Excel parse error: {str(e)}")

# ========== API Endpoints ==========

class RFQResponse(BaseModel):
    id: int
    city: Optional[str]
    total_budget: Optional[str]
    created_at: str
    updated_at: str
    items_count: int
    quotes_count: int
    status: str

class ItemResponse(BaseModel):
    id: int
    rfq_id: int
    name: str
    tech_code: Optional[str]
    gost: Optional[str]
    unit: Optional[str]
    qty: float
    deadline: Optional[str]
    location: Optional[str]
    material: Optional[str]
    suppliers_count: int
    quotes_count: int
    winner: Optional[dict] = None

class WinnerResponse(BaseModel):
    id: int
    quote_id: int
    final_price: float
    margin_pct: float
    chosen_at: str

class SelectWinnerRequest(BaseModel):
    quote_id: int
    margin_pct: float = 0.0

class QuoteResponse(BaseModel):
    id: int
    supplier_id: int
    supplier_name: Optional[str]
    supplier_email: str
    price: float
    delivery_days: Optional[int]
    pickup_address: Optional[str]
    source: str
    received_at: str

class SupplierResponse(BaseModel):
    id: int
    name: Optional[str]
    domain: str
    email: str
    is_valid: bool
    status: str

class ThreadResponse(BaseModel):
    id: int
    direction: str
    body: str
    created_at: str

class ProfileUpdate(BaseModel):
    company_name: Optional[str] = None
    inn: Optional[str] = None
    kpp: Optional[str] = None
    address: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    contact_person: Optional[str] = None

class CreateRFQResponse(BaseModel):
    rfq_id: int
    items_count: int
    message: str

@app.post("/api/rfqs", response_model=CreateRFQResponse)
async def create_rfq(
    file: UploadFile = File(...),
    vk_user_id: int = Depends(get_current_user)
):
    """Создать заявку из Excel файла"""
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Only .xlsx files supported")
    
    content = await file.read()
    
    try:
        items, city, budget = parse_excel_simple(content)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel: {str(e)}")
    
    if len(items) > 100:
        raise HTTPException(status_code=400, detail="Maximum 100 items allowed")
    
    if len(items) == 0:
        raise HTTPException(status_code=400, detail="No items found in Excel")
    
    async with db_pool.acquire() as conn:
        async with conn.transaction():
            # Создаем RFQ
            rfq_row = await conn.fetchrow(
                """
                INSERT INTO rfq_requests (city, total_budget, created_by_vk, created_at, updated_at)
                VALUES ($1, $2, $3, NOW(), NOW())
                RETURNING id
                """,
                city,
                str(budget) if budget else None,
                vk_user_id,
            )
            rfq_id = rfq_row["id"]
            
            # Создаем позиции
            for item in items:
                await conn.execute(
                    """
                    INSERT INTO items (
                        rfq_id, name, tech_code, gost, unit, qty, 
                        deadline, location, material, created_at, updated_at
                    ) VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, NOW(), NOW())
                    """,
                    rfq_id,
                    item["name"],
                    item.get("tech_code"),
                    item.get("gost"),
                    item.get("unit", "шт"),
                    item["qty"],
                    item.get("deadline"),
                    item.get("location") or city,
                    item.get("material"),
                )
            
            # Публикуем в Redis
            await redis_client.xadd("new_rfq", {
                "rfq_id": str(rfq_id),
                "items_count": str(len(items)),
                "created_by_vk": str(vk_user_id),
            })
    
    return CreateRFQResponse(
        rfq_id=rfq_id,
        items_count=len(items),
        message=f"Created RFQ #{rfq_id} with {len(items)} items"
    )

@app.get("/api/rfqs", response_model=List[RFQResponse])
async def get_rfqs(vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        rows = await conn.fetch("""
            SELECT 
                r.*,
                COUNT(DISTINCT i.id) as items_count,
                COUNT(DISTINCT q.id) as quotes_count,
                CASE 
                    WHEN EXISTS (SELECT 1 FROM winners w JOIN items wi ON w.item_id = wi.id WHERE wi.rfq_id = r.id)
                    THEN 'completed'
                    ELSE 'active'
                END as status
            FROM rfq_requests r
            LEFT JOIN items i ON i.rfq_id = r.id
            LEFT JOIN quotes q ON q.item_id = i.id
            WHERE r.created_by_vk = $1
            GROUP BY r.id
            ORDER BY r.created_at DESC
        """, vk_user_id)
        
        return [
            RFQResponse(
                id=row["id"],
                city=row["city"],
                total_budget=row["total_budget"],
                created_at=row["created_at"].isoformat(),
                updated_at=row["updated_at"].isoformat(),
                items_count=row["items_count"],
                quotes_count=row["quotes_count"],
                status=row["status"],
            )
            for row in rows
        ]

@app.get("/api/rfqs/{rfq_id}", response_model=RFQResponse)
async def get_rfq(rfq_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        row = await conn.fetchrow("""
            SELECT 
                r.*,
                COUNT(DISTINCT i.id) as items_count,
                COUNT(DISTINCT q.id) as quotes_count,
                CASE 
                    WHEN EXISTS (SELECT 1 FROM winners w JOIN items wi ON w.item_id = wi.id WHERE wi.rfq_id = r.id)
                    THEN 'completed'
                    ELSE 'active'
                END as status
            FROM rfq_requests r
            LEFT JOIN items i ON i.rfq_id = r.id
            LEFT JOIN quotes q ON q.item_id = i.id
            WHERE r.id = $1 AND r.created_by_vk = $2
            GROUP BY r.id
        """, rfq_id, vk_user_id)
        
        if not row:
            raise HTTPException(status_code=404, detail="RFQ not found")
        
        return RFQResponse(
            id=row["id"],
            city=row["city"],
            total_budget=row["total_budget"],
            created_at=row["created_at"].isoformat(),
            updated_at=row["updated_at"].isoformat(),
            items_count=row["items_count"],
            quotes_count=row["quotes_count"],
            status=row["status"],
        )

@app.get("/api/rfqs/{rfq_id}/items", response_model=List[ItemResponse])
async def get_items(rfq_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        rfq = await conn.fetchrow(
            "SELECT id FROM rfq_requests WHERE id = $1 AND created_by_vk = $2",
            rfq_id, vk_user_id
        )
        if not rfq:
            raise HTTPException(status_code=404, detail="RFQ not found")
        
        rows = await conn.fetch("""
            SELECT 
                i.*,
                COUNT(DISTINCT rs.id) as suppliers_count,
                COUNT(DISTINCT q.id) as quotes_count
            FROM items i
            LEFT JOIN rfq_sends rs ON rs.item_id = i.id
            LEFT JOIN quotes q ON q.item_id = i.id
            WHERE i.rfq_id = $1
            GROUP BY i.id
            ORDER BY i.id
        """, rfq_id)
        
        return [
            ItemResponse(
                id=row["id"],
                rfq_id=row["rfq_id"],
                name=row["name"],
                tech_code=row["tech_code"],
                gost=row["gost"],
                unit=row["unit"],
                qty=float(row["qty"]) if row["qty"] else 0,
                deadline=row["deadline"].isoformat() if row["deadline"] else None,
                location=row["location"],
                material=row["material"],
                suppliers_count=row["suppliers_count"],
                quotes_count=row["quotes_count"],
            )
            for row in rows
        ]

@app.get("/api/items/{item_id}", response_model=ItemResponse)
async def get_item(item_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        row = await conn.fetchrow("""
            SELECT 
                i.*,
                COUNT(DISTINCT rs.id) as suppliers_count,
                COUNT(DISTINCT q.id) as quotes_count
            FROM items i
            JOIN rfq_requests r ON r.id = i.rfq_id
            LEFT JOIN rfq_sends rs ON rs.item_id = i.id
            LEFT JOIN quotes q ON q.item_id = i.id
            WHERE i.id = $1 AND r.created_by_vk = $2
            GROUP BY i.id
        """, item_id, vk_user_id)
        
        if not row:
            raise HTTPException(status_code=404, detail="Item not found")
        
        # Fetch winner info
        winner_row = await conn.fetchrow("""
            SELECT w.* FROM winners w
            WHERE w.item_id = $1
        """, item_id)
        
        winner = None
        if winner_row:
            winner = {
                "id": winner_row["id"],
                "quote_id": winner_row["quote_id"],
                "final_price": float(winner_row["final_price"]),
                "margin_pct": float(winner_row["margin_pct"]),
                "chosen_at": winner_row["chosen_at"].isoformat() if winner_row["chosen_at"] else None,
            }
        
        return ItemResponse(
            id=row["id"],
            rfq_id=row["rfq_id"],
            name=row["name"],
            tech_code=row["tech_code"],
            gost=row["gost"],
            unit=row["unit"],
            qty=float(row["qty"]) if row["qty"] else 0,
            deadline=row["deadline"].isoformat() if row["deadline"] else None,
            location=row["location"],
            material=row["material"],
            suppliers_count=row["suppliers_count"],
            quotes_count=row["quotes_count"],
            winner=winner,
        )

@app.post("/api/items/{item_id}/winner", response_model=WinnerResponse)
async def select_winner(
    item_id: int, 
    request: SelectWinnerRequest,
    vk_user_id: int = Depends(get_current_user)
):
    async with db_pool.acquire() as conn:
        # Check access
        access = await conn.fetchrow("""
            SELECT 1 FROM items i
            JOIN rfq_requests r ON r.id = i.rfq_id
            WHERE i.id = $1 AND r.created_by_vk = $2
        """, item_id, vk_user_id)
        if not access:
            raise HTTPException(status_code=404, detail="Item not found")
        
        # Verify quote exists and belongs to this item
        quote = await conn.fetchrow("""
            SELECT price FROM quotes
            WHERE id = $1 AND item_id = $2
        """, request.quote_id, item_id)
        if not quote:
            raise HTTPException(status_code=400, detail="Quote not found for this item")
        
        # Insert or update winner
        row = await conn.fetchrow("""
            INSERT INTO winners (item_id, quote_id, final_price, margin_pct)
            VALUES ($1, $2, $3, $4)
            ON CONFLICT (item_id) DO UPDATE SET
                quote_id = EXCLUDED.quote_id,
                final_price = EXCLUDED.final_price,
                margin_pct = EXCLUDED.margin_pct,
                chosen_at = NOW()
            RETURNING *
        """, item_id, request.quote_id, quote["price"], request.margin_pct)
        
        return WinnerResponse(
            id=row["id"],
            quote_id=row["quote_id"],
            final_price=float(row["final_price"]),
            margin_pct=float(row["margin_pct"]),
            chosen_at=row["chosen_at"].isoformat() if row["chosen_at"] else None,
        )

@app.get("/api/items/{item_id}/quotes", response_model=List[QuoteResponse])
async def get_quotes(item_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        access = await conn.fetchrow("""
            SELECT 1 FROM items i
            JOIN rfq_requests r ON r.id = i.rfq_id
            WHERE i.id = $1 AND r.created_by_vk = $2
        """, item_id, vk_user_id)
        if not access:
            raise HTTPException(status_code=404, detail="Item not found")
        
        rows = await conn.fetch("""
            SELECT 
                q.*,
                s.name as supplier_name
            FROM quotes q
            LEFT JOIN suppliers s ON s.id = q.supplier_id
            WHERE q.item_id = $1
            ORDER BY q.price ASC
        """, item_id)
        
        return [
            QuoteResponse(
                id=row["id"],
                supplier_id=row["supplier_id"],
                supplier_name=row["supplier_name"],
                supplier_email=row["supplier_email"],
                price=float(row["price"]),
                delivery_days=row["delivery_days"],
                pickup_address=row["pickup_address"],
                source=row["source"],
                received_at=row["received_at"].isoformat() if row["received_at"] else None,
            )
            for row in rows
        ]

@app.get("/api/items/{item_id}/suppliers", response_model=List[SupplierResponse])
async def get_suppliers(item_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        access = await conn.fetchrow("""
            SELECT 1 FROM items i
            JOIN rfq_requests r ON r.id = i.rfq_id
            WHERE i.id = $1 AND r.created_by_vk = $2
        """, item_id, vk_user_id)
        if not access:
            raise HTTPException(status_code=404, detail="Item not found")
        
        rows = await conn.fetch("""
            SELECT DISTINCT
                s.id,
                s.name,
                s.domain,
                se.email,
                se.is_valid,
                COALESCE(
                    CASE 
                        WHEN o.status = 'unavailable' THEN 'unavailable'
                        WHEN st.direction = 'inbound' THEN 'replied'
                        WHEN rs.id IS NOT NULL THEN 'sent'
                        ELSE 'pending'
                    END,
                    'pending'
                ) as status
            FROM suppliers s
            JOIN supplier_emails se ON se.supplier_id = s.id
            JOIN rfq_sends rs ON rs.supplier_id = s.id AND rs.item_id = $1
            LEFT JOIN supplier_threads st ON st.rfq_send_id = rs.id AND st.direction = 'inbound'
            LEFT JOIN orphan_emails o ON o.from_email = se.email
            WHERE rs.item_id = $1
            ORDER BY s.id
        """, item_id)
        
        return [
            SupplierResponse(
                id=row["id"],
                name=row["name"],
                domain=row["domain"],
                email=row["email"],
                is_valid=row["is_valid"],
                status=row["status"],
            )
            for row in rows
        ]

@app.get("/api/items/{item_id}/threads", response_model=List[ThreadResponse])
async def get_threads(item_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        access = await conn.fetchrow("""
            SELECT 1 FROM items i
            JOIN rfq_requests r ON r.id = i.rfq_id
            WHERE i.id = $1 AND r.created_by_vk = $2
        """, item_id, vk_user_id)
        if not access:
            raise HTTPException(status_code=404, detail="Item not found")
        
        rows = await conn.fetch("""
            SELECT st.*
            FROM supplier_threads st
            JOIN rfq_sends rs ON rs.id = st.rfq_send_id
            WHERE rs.item_id = $1
            ORDER BY st.created_at DESC
            LIMIT 50
        """, item_id)
        
        return [
            ThreadResponse(
                id=row["id"],
                direction=row["direction"],
                body=row["body"],
                created_at=row["created_at"].isoformat(),
            )
            for row in rows
        ]

@app.get("/api/profile")
async def get_profile(vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        row = await conn.fetchrow("""
            SELECT up.*
            FROM user_profiles up
            WHERE up.vk_id = $1
        """, vk_user_id)
        
        if not row:
            return {}
        
        return dict(row)

@app.patch("/api/profile")
async def update_profile(data: ProfileUpdate, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        await conn.execute("""
            INSERT INTO user_profiles (vk_id, company_name, inn, kpp, address, phone, email, contact_person, updated_at)
            VALUES ($1, $2, $3, $4, $5, $6, $7, $8, NOW())
            ON CONFLICT (vk_id) DO UPDATE SET
                company_name = EXCLUDED.company_name,
                inn = EXCLUDED.inn,
                kpp = EXCLUDED.kpp,
                address = EXCLUDED.address,
                phone = EXCLUDED.phone,
                email = EXCLUDED.email,
                contact_person = EXCLUDED.contact_person,
                updated_at = NOW()
        """, 
            vk_user_id,
            data.company_name,
            data.inn,
            data.kpp,
            data.address,
            data.phone,
            data.email,
            data.contact_person,
        )
        
        return {"status": "ok"}

@app.get("/api/reports/{rfq_id}/status")
async def report_status(rfq_id: int, vk_user_id: int = Depends(get_current_user)):
    async with db_pool.acquire() as conn:
        access = await conn.fetchrow(
            "SELECT 1 FROM rfq_requests WHERE id = $1 AND created_by_vk = $2",
            rfq_id, vk_user_id
        )
        if not access:
            raise HTTPException(status_code=404, detail="RFQ not found")
        
        import os
        report_path = f"/app/reports/report_{rfq_id}.xlsx"
        exists = os.path.exists(report_path)
        
        return {"ready": exists, "rfq_id": rfq_id}

@app.get("/health")
async def health():
    return {"status": "ok", "version": "1.1.0"}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8002)
