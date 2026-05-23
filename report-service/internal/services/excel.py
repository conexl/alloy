"""Excel report generator."""

import os
from typing import List, Dict, Optional
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO

from ..domain.models import Item


class ExcelGenerator:
    """Excel report generator."""

    def __init__(self, report_dir: str):
        self.report_dir = report_dir
        os.makedirs(report_dir, exist_ok=True)

    def generate(
        self, rfq_id: int, items: List[Dict], profile: Optional[Dict] = None
    ) -> bytes:
        """Generate Excel report and return as bytes."""
        if not items:
            raise ValueError("No data for report")

        df = pd.DataFrame(items)

        column_mapping = {
            "num": "№",
            "name": "Наименование МТР",
            "tech_code": "Техобозначение",
            "material": "Материал",
            "gost": "ГОСТ/ТУ",
            "tt_ol_ref": "ТТ-ОЛ",
            "unit": "Ед.изм",
            "qty": "Кол-во",
            "deadline": "Срок поставки",
            "location": "Место назначения",
            "rkd_ref": "РКД",
            "payment_terms": "Условия оплаты",
            "final_price": "Цена без НДС",
            "supplier_email": "Email поставщика",
            "delivery_days": "Срок (дней)",
            "pickup_address": "Адрес самовывоза",
            "supplier_name": "Поставщик",
        }
        df = df.rename(
            columns={k: v for k, v in column_mapping.items() if k in df.columns}
        )

        for col in [
            "№",
            "Наименование МТР",
            "Техобозначение",
            "Материал",
            "ГОСТ/ТУ",
            "ТТ-ОЛ",
            "Ед.изм",
            "Кол-во",
            "Срок поставки",
            "Место назначения",
            "РКД",
            "Условия оплаты",
            "Цена без НДС",
            "Email поставщика",
            "Срок (дней)",
            "Адрес самовывоза",
            "Поставщик",
        ]:
            if col not in df.columns:
                df[col] = None

        output_columns = [
            "№",
            "Наименование МТР",
            "Техобозначение",
            "Материал",
            "ГОСТ/ТУ",
            "ТТ-ОЛ",
            "Ед.изм",
            "Кол-во",
            "Срок поставки",
            "Место назначения",
            "РКД",
            "Условия оплаты",
            "Цена без НДС",
            "Email поставщика",
            "Срок (дней)",
            "Адрес самовывоза",
            "Поставщик",
        ]
        df = df[output_columns]

        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=f"КП #{rfq_id}")
            worksheet = writer.sheets[f"КП #{rfq_id}"]

            header_fill = PatternFill(
                start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
            )
            header_font = Font(bold=True)

            if profile:
                profile_fields = [
                    ("Организация", "company_name"),
                    ("ИНН", "inn"),
                    ("КПП", "kpp"),
                    ("Адрес", "address"),
                    ("Телефон", "phone"),
                    ("Email", "email"),
                    ("Контактное лицо", "contact_person"),
                ]
                worksheet.insert_rows(1, 10)
                bold_font = Font(bold=True)
                for i, (label, key) in enumerate(profile_fields):
                    value = profile.get(key, "") or ""
                    cell_label = worksheet.cell(row=i + 1, column=1)
                    cell_label.value = label
                    cell_label.font = bold_font
                    cell_value = worksheet.cell(row=i + 1, column=2)
                    cell_value.value = value
                worksheet.cell(row=9, column=1).value = ""
                data_start_row = 10
            else:
                data_start_row = 1

            header_row = data_start_row
            for col_num, column_title in enumerate(df.columns, 1):
                cell = worksheet.cell(row=header_row, column=col_num)
                cell.fill = header_fill
                cell.font = header_font

            green_fill = PatternFill(
                start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"
            )
            yellow_fill = PatternFill(
                start_color="FFF9C4", end_color="FFF9C4", fill_type="solid"
            )
            for row_num in range(header_row + 1, header_row + len(df) + 1):
                final_price = worksheet.cell(
                    row=row_num, column=output_columns.index("Цена без НДС") + 1
                ).value
                if final_price is not None:
                    for col_num in range(1, len(df.columns) + 1):
                        worksheet.cell(row=row_num, column=col_num).fill = green_fill
                else:
                    for col_num in range(1, len(df.columns) + 1):
                        worksheet.cell(row=row_num, column=col_num).fill = yellow_fill

            worksheet.freeze_panes = f"A{data_start_row + 1}"

            for column in df.columns:
                column_letter = get_column_letter(output_columns.index(column) + 1)
                max_length = max(df[column].astype(str).map(len).max(), len(column))
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width

            total_row = header_row + len(df) + 1
            final_price_col = output_columns.index("Цена без НДС") + 1
            final_price_letter = get_column_letter(final_price_col)
            worksheet.cell(
                row=total_row, column=final_price_col
            ).value = f"=SUM({final_price_letter}{header_row + 1}:{final_price_letter}{total_row - 1})"
            worksheet.cell(row=total_row, column=final_price_col).font = Font(bold=True)

        output.seek(0)
        return output.read()

    def get_file_path(self, rfq_id: int) -> str:
        """Get file path for RFQ report."""
        return os.path.join(self.report_dir, f"rfq_{rfq_id}.xlsx")

    def save(self, rfq_id: int, content: bytes) -> str:
        """Save report to file and return path."""
        path = self.get_file_path(rfq_id)
        with open(path, "wb") as f:
            f.write(content)
        return path

    def file_exists(self, rfq_id: int) -> bool:
        """Check if report file exists."""
        return os.path.exists(self.get_file_path(rfq_id))
